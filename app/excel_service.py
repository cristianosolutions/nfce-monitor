from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def exportar_excel(registros, caminho_saida):
    caminho_saida = Path(caminho_saida)
    wb = Workbook()
    ws = wb.active
    ws.title = "NFC-e"

    colunas = [
        ("Número", "numero"), ("Série", "serie"), ("Emissão", "data_emissao"),
        ("CNPJ", "cnpj_formatado"), ("Emitente", "emitente"),
        ("Fantasia", "fantasia"), ("Valor", "valor"),
        ("Status XML", "status"), ("Situação SEFAZ", "situacao_sefaz"),
        ("cStat SEFAZ", "cstat_sefaz"), ("Motivo SEFAZ", "motivo_sefaz"),
        ("Protocolo SEFAZ", "protocolo_sefaz"),
        ("Data consulta SEFAZ", "data_consulta_sefaz"),
        ("Chave", "chave"), ("Arquivo", "arquivo"),
    ]

    ws.append([titulo for titulo, _ in colunas])
    for celula in ws[1]:
        celula.font = Font(bold=True)

    for registro in registros:
        ws.append([registro.get(campo, "") for _, campo in colunas])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, (titulo, _) in enumerate(colunas, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(
            12, min(45, len(titulo) + 6)
        )

    wb.save(caminho_saida)
    return str(caminho_saida)
